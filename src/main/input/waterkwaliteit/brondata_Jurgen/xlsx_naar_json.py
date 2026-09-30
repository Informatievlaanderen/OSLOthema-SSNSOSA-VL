"""Zet elk .xlsx-bestand in deze map om naar een subdirectory met één JSON-bestand per tab.

Gebruik (vanuit deze map):
    uv run --with openpyxl python xlsx_naar_json.py

Conventies:
- Subdirectory = bestandsnaam zonder extensie, JSON-bestand = tabnaam (spaties -> '_').
- Elke tab wordt een lijst van objecten; de sleutels zijn de kolomkoppen.
- Volledig lege rijen worden overgeslagen, lege cellen worden null.
- Een tab met een tweeregelige kop (groepsrij + kolomrij) gebruikt de tweede rij als sleutels.
- Een tab zonder kop krijgt sleutels 'kolom_A', 'kolom_B', ...
- Datums worden ISO 8601 ('2019-01-28' of '2019-01-28T13:01:00'), tijden 'HH:MM:SS'.
"""
import datetime
import glob
import json
import os

import openpyxl
from openpyxl.utils import get_column_letter


def waarde(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    return v


def is_kop(rij):
    """Een koprij bevat enkel tekst (of lege cellen) en minstens één tekstcel."""
    return any(isinstance(c, str) for c in rij) and all(c is None or isinstance(c, str) for c in rij)


def kop_en_data(rijen):
    if len(rijen) > 1 and is_kop(rijen[0]) and is_kop(rijen[1]) and None in rijen[0] and None not in rijen[1]:
        return list(rijen[1]), rijen[2:]
    if is_kop(rijen[0]) and None not in rijen[0]:
        return list(rijen[0]), rijen[1:]
    return [f"kolom_{get_column_letter(i + 1)}" for i in range(len(rijen[0]))], rijen


def zet_om(pad):
    doel = os.path.splitext(os.path.basename(pad))[0].replace(" ", "_")
    os.makedirs(doel, exist_ok=True)
    wb = openpyxl.load_workbook(pad, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rijen = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rijen:
            continue
        kop, data = kop_en_data(rijen)
        records = [{k: waarde(v) for k, v in zip(kop, r)} for r in data]
        uit = os.path.join(doel, ws.title.replace(" ", "_") + ".json")
        with open(uit, "w", encoding="utf-8") as f:
            json.dump(records, f, ensure_ascii=False, indent=2)
        print(f"{uit}: {len(records)} records, {len(kop)} kolommen")


if __name__ == "__main__":
    for pad in sorted(glob.glob("*.xlsx")):
        zet_om(pad)
